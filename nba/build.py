"""Build 2026-27-NBA-Roto-Values.xlsx: FantasyPros consensus projections valued
for a Yahoo public prize roto league (12 teams, 9-cat, 82 GP per slot).

The workbook is formula-driven, like the NHL projections sheet: change a cell
on Settings and every value, rank and DIFF recalculates.

    .venv/bin/python nba/build.py            # build from the saved pages in nba/data
    python3 nba/sources.py --refresh         # re-download the pages first

The chain (same in the formulas and in value_players() below):
  0. Minutes & Health: games = your GP, else the projected GP capped at the
     games left once he's back from injury (82 − team games before ESPN's
     return date). Minutes per game = your MPG, else the projected MPG.
  1. Projected totals = per-minute rate × MPG × games × boost (games = Flat GP
     when "Games Played Projection" is No). FG% and FT% carry over; attempts
     scale with minutes.
  2. Percentages become impact: FGA × (FG% − pool FG%), same for FT.
  3. Pass 1 z-scores every player against all projected players. The top
     "Value Pool" players by that score form the pool.
  4. Pass 2 z-scores everyone against the pool's mean and stdev. TO counts
     against you. VALUE = Σ weight × z.
  5. VORP = VALUE − VALUE of the first undrafted player (teams × roster + 1).
     RK = rank of VORP. PRK = rank among players with the same first position.
"""
import datetime
import statistics
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, str(Path(__file__).parent))
import sources  # noqa: E402

OUT = Path(__file__).parent / '2026-27-NBA-Roto-Values.xlsx'

# Yahoo public prize roto defaults (basketball.fantasysports.yahoo.com/nba/express_settings?type=roto).
CATEGORIES = ['PTS', 'REB', 'AST', 'STL', 'BLK', '3PM', 'FG%', 'FT%', 'TO']
SLOTS = [('PG', 1), ('SG', 1), ('G', 1), ('SF', 1), ('PF', 1), ('F', 1), ('C', 2), ('Util', 2), ('BN', 3), ('IL', 3)]
BOOSTS = [('+ + +', 0.15), ('+ +', 0.10), ('+', 0.05), ('-', -0.05), ('- -', -0.10), ('- - -', -0.15)]
DEFAULTS = {'teams': 12, 'gpProjection': 'Yes', 'flatGp': 75, 'boosts': 'ON'}

# Free-throw rate (FTA per FGA) is shrunk toward the position median by this
# many attempts, so a 40-shot season doesn't set a player's rate.
FTR_PRIOR_FGA = 200
SEASON_GAMES = 82  # bbref lists 80 per team; the other 2 are NBA Cup games not yet scheduled


# ---------------------------------------------------------------- data

def free_throw_rates(players, history):
    by_pos = {}
    for k, h in history.items():
        if h['fga'] >= 500 and h.get('pos'):
            by_pos.setdefault(h['pos'], []).append(h['fta'] / h['fga'])
    league = statistics.median(r for rs in by_pos.values() for r in rs)
    medians = {pos: statistics.median(rs) for pos, rs in by_pos.items()}
    for p in players:
        prior = medians.get(p['pos'][0], league) if p['pos'] else league
        h = history.get(sources.key(p['name']))
        if h and h['fga'] > 0:
            p['ftr'] = (h['fta'] + FTR_PRIOR_FGA * prior) / (h['fga'] + FTR_PRIOR_FGA)
            p['ftrSource'] = f"2025-26 ({h['fga']:.0f} FGA)"
        else:
            p['ftr'] = prior
            p['ftrSource'] = f"{p['pos'][0] if p['pos'] else 'league'} median"


def estimate_attempts(p):
    """PTS = 2·FGM + 3PM + FTM, FGM = FG%·FGA, FTA = FTr·FGA  ⇒  solve for FGA."""
    fga = (p['pts'] - p['tpm']) / (2 * p['fgPct'] + p['ftPct'] * p['ftr'])
    p['fga'], p['fta'] = round(fga, 1), round(p['ftr'] * fga, 1)


def load_players():
    players = sources.projections()
    adp = sources.adp()
    history = sources.last_season()
    # Position medians come from projected players' primary positions.
    for p in players:
        h = history.get(sources.key(p['name']))
        if h and p['pos']:
            h['pos'] = p['pos'][0]
    free_throw_rates(players, history)
    schedule, team_names = sources.schedule()
    hurt = sources.injuries(team_names)
    for p in players:
        estimate_attempts(p)
        a = adp.get(sources.key(p['name']), {})
        p['adpYahoo'], p['adpEspn'] = a.get('yahoo'), a.get('espn')
        h = history.get(sources.key(p['name']))
        p['gp25'], p['mpg25'] = (h['gp'], h['min'] / h['gp']) if h and h['gp'] else (None, None)
        i = hurt.get(sources.key(p['name']), {})
        p['status'], p['returnDate'] = i.get('status', ''), i.get('returnDate')
        p['injuryNote'] = f"{i['updated']}: {i['comment']}" if i else ''
    return players, schedule


def games_out(p, schedule):
    if not p.get('returnDate'):
        return 0
    return sum(1 for day, team in schedule if team == p['team'] and day < p['returnDate'])


def minutes_and_games(p, schedule, your_gp=None, your_mpg=None):
    gp = your_gp if your_gp is not None else min(p['gp'], SEASON_GAMES - games_out(p, schedule))
    mpg = your_mpg if your_mpg is not None else round(p['min'] / p['gp'], 1)  # as the sheet shows it
    return gp, mpg


# ---------------------------------------------------------------- model (mirror of the formulas)

def drafted(teams):
    return teams * sum(n for slot, n in SLOTS if slot != 'IL')


def value_players(players, schedule, settings=DEFAULTS, weights=None, adj=None, minutes=None):
    """minutes: {name: (your_gp or None, your_mpg or None, return date override or None)}"""
    weights = weights or {c: 1 for c in CATEGORIES}
    adj = adj or {}
    minutes = minutes or {}
    boost = dict(BOOSTS)
    rows = []
    for p in players:
        your_gp, your_mpg, ret = minutes.get(p['name'], (None, None, None))
        q = {**p, 'returnDate': ret} if ret is not None else p
        gp_used, mpg = minutes_and_games(q, schedule, your_gp, your_mpg)
        games = gp_used if settings['gpProjection'] == 'Yes' else settings['flatGp']
        m = 1 + (boost.get(adj.get(p['name'], ''), 0) if settings['boosts'] == 'ON' else 0)
        f = mpg * games / p['min'] * m
        r = {'name': p['name'], 'pos': p['pos'][0] if p['pos'] else '',
             'PTS': p['pts'] * f, 'REB': p['reb'] * f, 'AST': p['ast'] * f, 'STL': p['stl'] * f,
             'BLK': p['blk'] * f, '3PM': p['tpm'] * f, 'TO': p['to'] * f,
             'fga': p['fga'] * f, 'fta': p['fta'] * f, 'fgPct': p['fgPct'], 'ftPct': p['ftPct']}
        rows.append(r)

    def zscores(members, weigh):
        fg = sum(r['fga'] * r['fgPct'] for r in members) / sum(r['fga'] for r in members)
        ft = sum(r['fta'] * r['ftPct'] for r in members) / sum(r['fta'] for r in members)
        for r in rows:
            r['FG%'] = r['fga'] * (r['fgPct'] - fg)
            r['FT%'] = r['fta'] * (r['ftPct'] - ft)
        stats = {c: (statistics.mean(r[c] for r in members), statistics.stdev(r[c] for r in members))
                 for c in CATEGORIES}
        for r in rows:
            r['z'] = {c: (r[c] - mu) / sd * (-1 if c == 'TO' else 1) for c, (mu, sd) in stats.items()}
            weigh(r, sum(weights[c] * r['z'][c] for c in CATEGORIES))

    zscores(rows, lambda r, v: r.__setitem__('base', v))
    pool_size = drafted(settings['teams'])
    cut = sorted((r['base'] for r in rows), reverse=True)[pool_size - 1]
    pool = [r for r in rows if r['base'] >= cut]
    zscores(pool, lambda r, v: r.__setitem__('value', v))
    replacement = sorted((r['value'] for r in rows), reverse=True)[pool_size]
    for r in rows:
        r['vorp'] = r['value'] - replacement
    for r in rows:
        r['rank'] = 1 + sum(o['vorp'] > r['vorp'] for o in rows)
        r['posRank'] = 1 + sum(o['vorp'] > r['vorp'] for o in rows if o['pos'] == r['pos'])
    return {r['name']: r for r in rows}


# ---------------------------------------------------------------- workbook

HEADER = Font(bold=True, color='FFFFFF')
HEADER_FILL = PatternFill('solid', fgColor='1F3864')
INPUT_FILL = PatternFill('solid', fgColor='FFF2CC')
NOTE = Font(italic=True, color='595959')


def head(ws, row, labels, start=1):
    for i, text in enumerate(labels):
        if text is None:
            continue
        c = ws.cell(row=row, column=start + i, value=text)
        c.font, c.fill = HEADER, HEADER_FILL
        c.alignment = Alignment(horizontal='center')


def build_settings(wb):
    ws = wb.active
    ws.title = 'Settings'
    head(ws, 1, ['Category', 'Weight'])
    for i, cat in enumerate(CATEGORIES):
        ws.cell(row=2 + i, column=1, value=cat)
        ws.cell(row=2 + i, column=2, value=1).fill = INPUT_FILL
    ws['A12'] = '1 = counts, 0 = punt, or any weight in between.'
    ws['A13'] = 'TO is scored in reverse automatically.'
    ws['A12'].font = ws['A13'].font = NOTE

    head(ws, 1, ['League', None, 'Notes'], start=4)
    rows = [
        ('Number of Teams', DEFAULTS['teams'], 'Yahoo public prize leagues: 12'),
        ('Players Drafted', '=F2*SUM(I2:I10)', 'teams × (starters + bench); IL excluded'),
        ('Replacement Rank', '=F3+1', 'first player left on waivers'),
        ('Value Pool', '=F3', 'players the averages / stdevs are measured against'),
        (None, None, None),
        ('Games Played Projection', DEFAULTS['gpProjection'], '"Yes" = projected GP, "No" = Flat GP for everyone'),
        ('Flat GP', DEFAULTS['flatGp'], 'used when Games Played Projection is No'),
        (None, None, None),
        ('Boost or Bust', DEFAULTS['boosts'], '"ON" / "OFF" — your ADJ marks on The List'),
    ]
    for i, (label, value, note) in enumerate(rows):
        r = 2 + i
        if label:
            ws.cell(row=r, column=4, value=label)
            c = ws.cell(row=r, column=6, value=value)
            if not str(value).startswith('='):
                c.fill = INPUT_FILL
            ws.cell(row=r, column=7, value=note).font = NOTE
    for i, (mark, adj) in enumerate(BOOSTS):
        ws.cell(row=12 + i, column=4, value=mark)
        ws.cell(row=12 + i, column=6, value=adj).fill = INPUT_FILL
    ws['G12'] = 'Multiplies every counting stat and attempt; FG% / FT% are unchanged.'
    ws['G12'].font = NOTE

    head(ws, 1, ['Slot', 'Count'], start=8)
    for i, (slot, n) in enumerate(SLOTS):
        ws.cell(row=2 + i, column=8, value=slot)
        ws.cell(row=2 + i, column=9, value=n).fill = INPUT_FILL
    ws['H13'] = 'Max games played: 82 per slot'
    ws['H13'].font = NOTE

    head(ws, 1, ['DIFF (ADP − RK)', 'Meaning'], start=11)
    for i, (label, meaning, color) in enumerate([
        ('≥ +24', 'Steal (>2 rounds)', '63BE7B'), ('+12 to +23', 'Potential steal (1–2 rounds)', 'C6EFCE'),
        ('−12 to −23', 'Potential reach (1–2 rounds)', 'FFC7CE'), ('≤ −24', 'Reach (>2 rounds)', 'F8696B')]):
        ws.cell(row=2 + i, column=11, value=label).fill = PatternFill('solid', fgColor=color)
        ws.cell(row=2 + i, column=12, value=meaning)

    yes_no = DataValidation(type='list', formula1='"Yes,No"', allow_blank=False)
    on_off = DataValidation(type='list', formula1='"ON,OFF"', allow_blank=False)
    ws.add_data_validation(yes_no)
    ws.add_data_validation(on_off)
    yes_no.add('F7')
    on_off.add('F10')
    for col, width in {'A': 10, 'B': 8, 'C': 2, 'D': 24, 'E': 2, 'F': 8, 'G': 48, 'H': 8, 'I': 8, 'J': 2, 'K': 16, 'L': 28}.items():
        ws.column_dimensions[col].width = width


PD_COLS = ['Player', 'Team', 'Pos', 'GP', 'MIN', 'PTS', 'REB', 'AST', 'STL', 'BLK', '3PM', 'TO',
           'FG%', 'FT%', 'FGA (est.)', 'FTA (est.)', 'FTr source', 'Yahoo ADP', 'ESPN ADP']


def build_player_data(wb, players):
    ws = wb.create_sheet('Player Data')
    head(ws, 1, PD_COLS)
    for r, p in enumerate(players, start=2):
        values = [p['name'], p['team'], ','.join(p['pos']), p['gp'], p['min'], p['pts'], p['reb'], p['ast'],
                  p['stl'], p['blk'], p['tpm'], p['to'], p['fgPct'], p['ftPct'], p['fga'],
                  p['fta'], p['ftrSource'], p['adpYahoo'], p['adpEspn']]
        for c, v in enumerate(values, start=1):
            ws.cell(row=r, column=c, value=v)
        ws.cell(row=r, column=13).number_format = ws.cell(row=r, column=14).number_format = '.000'
    ws.freeze_panes = 'B2'
    ws.column_dimensions['A'].width = 26
    ws.column_dimensions['Q'].width = 20
    return ws


MH_COLS = ['Player', 'Team', 'Pos', 'Proj GP', 'Proj MPG', '25-26 GP', '25-26 MPG', 'MPG Δ', 'Status',
           'Return', 'Games Out', 'Max GP', 'Your GP', 'Your MPG', 'GP Used', 'MPG Used', 'Note']
MH = {name: get_column_letter(i + 1) for i, name in enumerate(MH_COLS)}
ROLE_CHANGE_MPG = 3  # highlight projected minutes this far from last season


def build_minutes(wb, players):
    """Minutes & Health: projected vs last-season minutes, injuries, and your overrides.
    Same row order as Player Data, so The List reads it with the same ROW index."""
    ws = wb.create_sheet('Minutes & Health', 1)
    head(ws, 1, MH_COLS)
    for r, p in enumerate(players, start=2):
        c = {name: f'{MH[name]}{r}' for name in MH}
        values = {
            'Player': p['name'], 'Team': p['team'], 'Pos': ','.join(p['pos']),
            'Proj GP': p['gp'], 'Proj MPG': round(p['min'] / p['gp'], 1),
            '25-26 GP': p['gp25'], '25-26 MPG': round(p['mpg25'], 1) if p['mpg25'] else None,
            'MPG Δ': f'=IF({c["25-26 MPG"]}="","",{c["Proj MPG"]}-{c["25-26 MPG"]})',
            'Status': p['status'] or None,
            'Return': p['returnDate'],
            'Games Out': f'=IF({c["Return"]}="",0,COUNTIFS(Schedule!$B:$B,{c["Team"]},Schedule!$A:$A,"<"&{c["Return"]}))',
            'Max GP': f'={SEASON_GAMES}-{c["Games Out"]}',
            'GP Used': f'=IF({c["Your GP"]}<>"",{c["Your GP"]},MIN({c["Proj GP"]},{c["Max GP"]}))',
            'MPG Used': f'=IF({c["Your MPG"]}<>"",{c["Your MPG"]},{c["Proj MPG"]})',
            'Note': p['injuryNote'] or None,
        }
        for name, v in values.items():
            ws[c[name]] = v
        for name in ('Your GP', 'Your MPG', 'Return'):
            ws[c[name]].fill = INPUT_FILL
        ws[c['Return']].number_format = 'mmm d'
        ws[c['MPG Δ']].number_format = '+0.0;-0.0;0.0'
    last = len(players) + 1
    delta = f"{MH['MPG Δ']}2:{MH['MPG Δ']}{last}"
    ws.conditional_formatting.add(delta, CellIsRule(operator='greaterThanOrEqual', formula=[str(ROLE_CHANGE_MPG)],
                                                    fill=PatternFill('solid', fgColor='C6EFCE')))
    ws.conditional_formatting.add(delta, CellIsRule(operator='lessThanOrEqual', formula=[str(-ROLE_CHANGE_MPG)],
                                                    fill=PatternFill('solid', fgColor='FFC7CE')))
    ws.freeze_panes = 'B2'
    ws.auto_filter.ref = f'A1:{MH["Note"]}{last}'
    ws.column_dimensions['A'].width = 26
    ws.column_dimensions[MH['Note']].width = 90
    ws.column_dimensions[MH['Status']].width = 12
    note = len(players) + 3
    ws[f'A{note}'] = (f'Yellow cells are yours. Your GP / Your MPG replace the projection for that player. '
                      f'Return comes from ESPN; Games Out counts his team\'s games before it, and GP is capped at '
                      f'{SEASON_GAMES} − Games Out. MPG Δ is green/red when projected minutes moved '
                      f'{ROLE_CHANGE_MPG}+ from last season.')
    ws[f'A{note}'].font = NOTE


def build_schedule(wb, schedule):
    ws = wb.create_sheet('Schedule')
    head(ws, 1, ['Date', 'Team'])
    for r, (day, team) in enumerate(schedule, start=2):
        ws.cell(row=r, column=1, value=day).number_format = 'yyyy-mm-dd'
        ws.cell(row=r, column=2, value=team)
    ws.column_dimensions['A'].width = 12


# The List layout. Columns are named here once; formulas use these letters.
LIST = ['RK', 'NAME', 'POS', 'PRK', 'TEAM', 'ADJ', 'VALUE', 'VORP', 'ADP', 'DIFF', None,
        'GP', 'MPG', 'PTS', 'REB', 'AST', 'STL', 'BLK', '3PM', 'TO', 'FG%', 'FT%', 'FGM', 'FGA', 'FTM', 'FTA', None,
        'zPTS', 'zREB', 'zAST', 'zSTL', 'zBLK', 'z3PM', 'zFG%', 'zFT%', 'zTO', None,
        'ROW', 'MULT', '1POS', 'FGimp1', 'FTimp1', 'BASE', 'POOL', 'FGimp2', 'FTimp2']
COL = {name: get_column_letter(i + 1) for i, name in enumerate(LIST) if name}
STAT_FOR_Z = {'PTS': 'PTS', 'REB': 'REB', 'AST': 'AST', 'STL': 'STL', 'BLK': 'BLK', '3PM': '3PM',
              'FG%': 'FGimp2', 'FT%': 'FTimp2', 'TO': 'TO'}
STAT_FOR_BASE = {**STAT_FOR_Z, 'FG%': 'FGimp1', 'FT%': 'FTimp1'}
PD = {name: get_column_letter(i + 1) for i, name in enumerate(PD_COLS)}


def build_pool(wb, last):
    """Means and stdevs: pass 1 over everyone, pass 2 over the value pool."""
    ws = wb.create_sheet('Pool')
    head(ws, 1, ['Category', 'All mean', 'All stdev', 'Pool mean', 'Pool stdev'])
    rng = lambda c: f"'The List'!${COL[c]}$2:${COL[c]}${last}"
    pool = rng('POOL')
    for i, cat in enumerate(CATEGORIES):
        r = 2 + i
        a, b = rng(STAT_FOR_BASE[cat]), rng(STAT_FOR_Z[cat])
        ws.cell(row=r, column=1, value=cat)
        ws.cell(row=r, column=2, value=f'=AVERAGE({a})')
        ws.cell(row=r, column=3, value=f'=STDEV({a})')
        ws.cell(row=r, column=4, value=f'=SUMPRODUCT({pool},{b})/SUM({pool})')
        ws.cell(row=r, column=5, value=f'=SQRT(SUMPRODUCT({pool},({b}-D{r})^2)/(SUM({pool})-1))')
    ws['A12'], ws['B12'], ws['D12'] = 'FG%', f"=SUM({rng('FGM')})/SUM({rng('FGA')})", \
        f"=SUMPRODUCT({pool},{rng('FGM')})/SUMPRODUCT({pool},{rng('FGA')})"
    ws['A13'], ws['B13'], ws['D13'] = 'FT%', f"=SUM({rng('FTM')})/SUM({rng('FTA')})", \
        f"=SUMPRODUCT({pool},{rng('FTM')})/SUMPRODUCT({pool},{rng('FTA')})"
    ws['A11'] = 'Shooting average (makes ÷ attempts)'
    ws['A11'].font = NOTE
    for row in ws.iter_rows(min_row=2, max_row=13, min_col=2, max_col=5):
        for c in row:
            c.number_format = '0.000'
    ws.column_dimensions['A'].width = 34
    return ws


def build_list(wb, players, valued):
    ws = wb.create_sheet('The List', 0)
    head(ws, 1, LIST)
    order = sorted(players, key=lambda p: valued[p['name']]['rank'])
    last = len(order) + 1
    span = lambda c: f'${COL[c]}$2:${COL[c]}${last}'
    pd_rows = {p['name']: i for i, p in enumerate(players, start=2)}

    for r, p in enumerate(order, start=2):
        c = {name: f'{COL[name]}{r}' for name in COL}
        pd = lambda field: f"INDEX('Player Data'!${PD[field]}:${PD[field]},{c['ROW']})"
        mh = lambda field: f"INDEX('Minutes & Health'!${MH[field]}:${MH[field]},{c['ROW']})"
        f = f"/{pd('MIN')}*{c['MPG']}*{c['GP']}*{c['MULT']}"
        formulas = {
            'ROW': f"=MATCH({c['NAME']},'Player Data'!$A:$A,0)",
            'NAME': p['name'],
            'POS': f"={pd('Pos')}",
            'TEAM': f"={pd('Team')}",
            '1POS': f'=IFERROR(LEFT({c["POS"]},FIND(",",{c["POS"]})-1),{c["POS"]})',
            'MULT': f"=1+IF(Settings!$F$10=\"ON\",IFERROR(VLOOKUP(TRIM({c['ADJ']}),Settings!$D$12:$F$17,3,FALSE),0),0)",
            'GP': f"=IF(Settings!$F$7=\"Yes\",{mh('GP Used')},Settings!$F$8)",
            'MPG': f"={mh('MPG Used')}",
            **{s: f"={pd(s)}{f}" for s in ['PTS', 'REB', 'AST', 'STL', 'BLK', '3PM', 'TO']},
            'FG%': f"={pd('FG%')}",
            'FT%': f"={pd('FT%')}",
            'FGA': f"={pd('FGA (est.)')}{f}",
            'FTA': f"={pd('FTA (est.)')}{f}",
            'FGM': f"={c['FGA']}*{c['FG%']}",
            'FTM': f"={c['FTA']}*{c['FT%']}",
            'FGimp1': f"={c['FGA']}*({c['FG%']}-Pool!$B$12)",
            'FTimp1': f"={c['FTA']}*({c['FT%']}-Pool!$B$13)",
            'FGimp2': f"={c['FGA']}*({c['FG%']}-Pool!$D$12)",
            'FTimp2': f"={c['FTA']}*({c['FT%']}-Pool!$D$13)",
            'BASE': '=' + '+'.join(
                f"Settings!$B${2 + i}*{'-' if cat == 'TO' else ''}({c[STAT_FOR_BASE[cat]]}-Pool!$B${2 + i})/Pool!$C${2 + i}"
                for i, cat in enumerate(CATEGORIES)),
            'POOL': f"=IF(RANK({c['BASE']},{span('BASE')})<=Settings!$F$5,1,0)",
            **{f'z{cat}': f"={'-' if cat == 'TO' else ''}({c[STAT_FOR_Z[cat]]}-Pool!$D${2 + i})/Pool!$E${2 + i}"
               for i, cat in enumerate(CATEGORIES)},
            'VALUE': '=' + '+'.join(f"Settings!$B${2 + i}*{c['z' + cat]}" for i, cat in enumerate(CATEGORIES)),
            'VORP': f"={c['VALUE']}-LARGE({span('VALUE')},Settings!$F$4)",
            'RK': f"=RANK({c['VORP']},{span('VORP')})",
            # Numeric compare: COUNTIFS(">"&x) rounds x to 15 digits as text.
            'PRK': f"=SUMPRODUCT(({span('1POS')}={c['1POS']})*({span('VORP')}>{c['VORP']}))+1",
            'ADP': f"=IF({pd('Yahoo ADP')}=\"\",\"\",{pd('Yahoo ADP')})",
            'DIFF': f"=IF({c['ADP']}=\"\",\"\",{c['ADP']}-{c['RK']})",
        }
        for name, formula in formulas.items():
            ws[c[name]] = formula
        ws[c['ADJ']].fill = INPUT_FILL

    fmt = {'VALUE': '0.00', 'VORP': '0.00', 'DIFF': '+0;-0;0', 'MPG': '0.0', 'GP': '0',
           'FG%': '.000', 'FT%': '.000', 'MULT': '0.00', 'BASE': '0.00',
           **{s: '0' for s in ['PTS', 'REB', 'AST', 'STL', 'BLK', '3PM', 'TO', 'FGM', 'FGA', 'FTM', 'FTA']},
           **{f'z{cat}': '0.00' for cat in CATEGORIES},
           **{s: '0.0' for s in ['FGimp1', 'FTimp1', 'FGimp2', 'FTimp2']}}
    for name, number_format in fmt.items():
        for row in ws[f'{COL[name]}2:{COL[name]}{last}']:
            row[0].number_format = number_format

    diff = f"{COL['DIFF']}2:{COL['DIFF']}{last}"
    for op, args, color in [('greaterThanOrEqual', ['24'], '63BE7B'), ('between', ['12', '23.99'], 'C6EFCE'),
                            ('between', ['-23.99', '-12'], 'FFC7CE'), ('lessThanOrEqual', ['-24'], 'F8696B')]:
        ws.conditional_formatting.add(diff, CellIsRule(operator=op, formula=args, fill=PatternFill('solid', fgColor=color)))
    zs = f"{COL['zPTS']}2:{COL['zTO']}{last}"
    ws.conditional_formatting.add(zs, ColorScaleRule(start_type='num', start_value=-2, start_color='F8696B',
                                                     mid_type='num', mid_value=0, mid_color='FFFFFF',
                                                     end_type='num', end_value=2, end_color='63BE7B'))

    ws.freeze_panes = 'C2'
    ws.auto_filter.ref = f'A1:{COL["zTO"]}{last}'
    ws.column_dimensions[COL['NAME']].width = 26
    for name in LIST:
        if name and name not in ('NAME',):
            ws.column_dimensions[COL[name]].width = 7
    ws.column_dimensions[COL['POS']].width = 10
    # Helper columns: grouped and collapsed so they stay out of the way.
    ws.column_dimensions.group(COL['ROW'], COL['FTimp2'], hidden=True)
    return last


def main():
    players, schedule = load_players()
    valued = value_players(players, schedule)
    wb = Workbook()
    build_settings(wb)
    last = build_list(wb, players, valued)
    build_minutes(wb, players)
    build_player_data(wb, players)
    build_pool(wb, last)
    build_schedule(wb, schedule)
    wb.calculation.fullCalcOnLoad = True
    wb.save(OUT)
    print(f'wrote {OUT.name}: {len(players)} players')
    for r in sorted(valued.values(), key=lambda r: r['rank'])[:15]:
        print(f"{r['rank']:>3} {r['name']:<28} {r['value']:6.2f} {r['vorp']:6.2f}")


if __name__ == '__main__':
    main()
