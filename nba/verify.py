"""Check an Excel-recalculated copy of the workbook against value_players().

    .venv/bin/python nba/verify.py <recalculated.xlsx>

Reads the Settings and ADJ marks the copy was calculated with, runs the Python
model under the same inputs, and compares every player's VALUE, VORP, RK and
PRK. Exits non-zero on any mismatch.
"""
import sys
from pathlib import Path

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).parent))
import build  # noqa: E402


def main(path):
    wb = load_workbook(path, data_only=True)
    s = wb['Settings']
    if s['F2'].value is None or wb['The List']['G2'].value is None:
        sys.exit('no cached values — recalculate and save the copy in Excel first')
    settings = {'teams': s['F2'].value, 'gpProjection': s['F7'].value, 'flatGp': s['F8'].value, 'boosts': s['F10'].value}
    weights = {s.cell(row=r, column=1).value: s.cell(row=r, column=2).value for r in range(2, 11)}
    slots = [(s.cell(row=r, column=8).value, s.cell(row=r, column=9).value) for r in range(2, 12)]
    assert slots == build.SLOTS, f'roster slots changed: {slots}'

    ws = wb['The List']
    col = {ws.cell(row=1, column=c).value: c for c in range(1, ws.max_column + 1)}
    sheet = {}
    for r in range(2, ws.max_row + 1):
        get = lambda name: ws.cell(row=r, column=col[name]).value
        sheet[get('NAME')] = {'adj': get('ADJ') or '', 'value': get('VALUE'), 'vorp': get('VORP'),
                              'rank': get('RK'), 'posRank': get('PRK')}

    players, schedule = build.load_players()
    original = {p['name']: p['returnDate'] for p in players}
    mh = wb['Minutes & Health']
    mcol = {mh.cell(row=1, column=c).value: c for c in range(1, mh.max_column + 1)}
    minutes = {}
    for r in range(2, len(players) + 2):
        get = lambda name: mh.cell(row=r, column=mcol[name]).value
        ret = get('Return')
        ret = ret.date() if hasattr(ret, 'date') else ret
        your_gp, your_mpg = get('Your GP'), get('Your MPG')
        if your_gp is not None or your_mpg is not None or ret != original[get('Player')]:
            minutes[get('Player')] = (your_gp, your_mpg, ret if ret != original[get('Player')] else None)

    valued = build.value_players(players, schedule, settings, weights,
                                 {n: v['adj'] for n, v in sheet.items() if v['adj']}, minutes)
    bad = []
    for name, x in sheet.items():
        y = valued[name]
        if abs(x['value'] - y['value']) > 1e-6 or abs(x['vorp'] - y['vorp']) > 1e-6 \
                or x['rank'] != y['rank'] or x['posRank'] != y['posRank']:
            bad.append((name, x, {k: y[k] for k in ('value', 'vorp', 'rank', 'posRank')}))
    print(f'{path}: {settings} weights={weights} minutes={minutes}')
    print(f'{len(sheet) - len(bad)}/{len(sheet)} players match')
    for b in bad[:10]:
        print('  MISMATCH', *b)
    sys.exit(1 if bad else 0)


if __name__ == '__main__':
    main(sys.argv[1])
